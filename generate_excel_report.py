import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report():
    # 1. Load exported data from live Supabase export
    with open('stalls_data_export.json', 'r', encoding='utf-8') as f:
        data = json.load(f)

    stalls = data['stalls']
    reviews = data['reviews']

    # Sort stalls numerically by stall_number
    stalls.sort(key=lambda s: int(''.join(filter(str.isdigit, str(s.get('stall_number', '0')))) or 0))

    # Pre-calculate stats for ordering the leaderboard cleanly
    stall_stats = {}
    for s in stalls:
        s_num = str(s.get('stall_number', '')).strip()
        s_id = str(s.get('id', '')).strip()
        num_clean = str(int(''.join(filter(str.isdigit, s_num)) or 0))
        
        matched_revs = [
            r for r in reviews
            if str(r.get('stall_id', '')).strip() in [s_id, s_num, f"stall_{s_num}", f"stall_{num_clean}", num_clean]
        ]
        
        total_rev = len(matched_revs)
        ratings = [r.get('rating', 0) for r in matched_revs if r.get('rating')]
        avg_rating = round(sum(ratings) / total_rev, 2) if total_rev > 0 else 0.0
        
        star_counts = {1: 0, 2: 0, 3: 0, 4: 0, 5: 0}
        tag_counts = {}
        for r in matched_revs:
            star = r.get('rating', 0)
            if star in star_counts:
                star_counts[star] += 1
            tags = r.get('review_text', '')
            if tags:
                for t in tags.split(','):
                    clean_t = t.strip()
                    if clean_t:
                        tag_counts[clean_t] = tag_counts.get(clean_t, 0) + 1

        top_tags = sorted(tag_counts.items(), key=lambda x: x[1], reverse=True)[:3]
        top_tags_str = ", ".join([f"{t[0]} ({t[1]})" for t in top_tags]) if top_tags else "—"

        stall_stats[s_id] = {
            'stall': s,
            'total_reviews': total_rev,
            'avg_rating': avg_rating,
            'stars': star_counts,
            'top_tags': top_tags_str,
            'reviews_list': matched_revs
        }

    # Rank stalls by avg_rating desc, then total_reviews desc, then stall_number asc
    ranked_stall_ids = sorted(
        stall_stats.keys(),
        key=lambda sid: (
            stall_stats[sid]['avg_rating'],
            stall_stats[sid]['total_reviews'],
            -int(''.join(filter(str.isdigit, str(stall_stats[sid]['stall'].get('stall_number', '0')))) or 0)
        ),
        reverse=True
    )

    wb = Workbook()
    
    # ----------------------------------------------------
    # STYLES & COLOR PALETTE
    # ----------------------------------------------------
    PRIMARY_COLOR = "8B0D1A"   # Adamas Crimson
    SECONDARY_COLOR = "1E293B" # Slate Navy
    ACCENT_GOLD = "D97706"     # Gold
    ACCENT_GREEN = "047857"    # Forest Green
    LIGHT_BG = "FAF8F5"       # Warm cream
    ZEBRA_BG = "F9FAFB"       # Soft alternating gray
    WHITE = "FFFFFF"
    BORDER_COLOR = "D1D5DB"
    HEADER_TEXT_COLOR = "FFFFFF"
    
    font_title = Font(name="Calibri", size=15, bold=True, color=WHITE)
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="F8FAFC")
    font_tbl_header = Font(name="Calibri", size=10, bold=True, color=HEADER_TEXT_COLOR)
    font_data = Font(name="Calibri", size=10)
    font_data_bold = Font(name="Calibri", size=10, bold=True)
    font_kpi_num = Font(name="Calibri", size=17, bold=True, color=PRIMARY_COLOR)
    font_kpi_lbl = Font(name="Calibri", size=9, bold=True, color="64748B")
    
    fill_header_primary = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    fill_header_secondary = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")
    fill_kpi_card = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_top1 = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Gold tint
    fill_top2 = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Silver tint
    fill_top3 = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid") # Bronze tint
    
    thin_side = Side(style='thin', color=BORDER_COLOR)
    medium_side = Side(style='medium', color="94A3B8")
    double_bottom = Side(style='double', color=PRIMARY_COLOR)
    
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=medium_side, bottom=medium_side)
    border_total_row = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_bottom)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    total_reviews_count = len(reviews)
    last_review_row = 3 + total_reviews_count

    # ----------------------------------------------------
    # SHEET 3: RAW VISITOR FEEDBACK FEED (Constructed First for clean cross-references)
    # ----------------------------------------------------
    ws3 = wb.active
    ws3.title = "All Visitor Reviews (1,178)"
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.merge_cells("A1:I1")
    ws3["A1"] = "ADAMAS UNIVERSITY E-CLUB — RAKHI STARTUP BAZAAR VISITOR RATINGS AUDIT FEED"
    ws3["A1"].font = font_title
    ws3["A1"].fill = fill_header_primary
    ws3["A1"].alignment = align_center
    ws3.row_dimensions[1].height = 26
    
    rev_headers = [
        ("Review ID", 12, align_center),
        ("Timestamp (UTC)", 20, align_center),
        ("Stall Code", 12, align_center),
        ("Stall #", 10, align_center),
        ("Stall Name", 28, align_left),
        ("Category", 24, align_left),
        ("Rating (1-5)", 13, align_center),
        ("Star Score Display", 18, align_center),
        ("Visitor Highlights & Standout Features", 52, align_left),
    ]
    
    h_row_rev = 3
    ws3.row_dimensions[h_row_rev].height = 24
    for col_idx, (h_title, col_width, col_align) in enumerate(rev_headers, 1):
        c = ws3.cell(row=h_row_rev, column=col_idx)
        c.value = h_title
        c.font = font_tbl_header
        c.fill = fill_header_secondary
        c.alignment = align_center
        c.border = border_header
        ws3.column_dimensions[get_column_letter(col_idx)].width = col_width

    # Sort reviews chronologically descending
    sorted_reviews = sorted(reviews, key=lambda r: r.get('created_at', ''), reverse=True)
    
    rev_data_row = 4
    for r_idx, r in enumerate(sorted_reviews, 1):
        sid = str(r.get('stall_id', ''))
        num_clean = str(int(''.join(filter(str.isdigit, sid)) or 0))
        matched_s = next(
            (s for s in stalls if s.get('id') in [sid, f"stall_{num_clean}"] or str(s.get('stall_number')) in [sid, num_clean]),
            None
        )
        
        stall_num_str = f"Stall #{str(matched_s.get('stall_number', '')).zfill(2)}" if matched_s else sid
        stall_code = str(matched_s.get('stall_number', '')).zfill(2) if matched_s else num_clean.zfill(2)
        stall_name_str = matched_s.get('name', '—') if matched_s else '—'
        stall_cat_str = matched_s.get('category', '—') if matched_s else '—'
        
        star_num = int(r.get('rating', 5))
        star_display = f"{'★' * star_num}{'☆' * (5 - star_num)} ({star_num}/5)"
        
        ts = r.get('created_at', '')
        ts_clean = ts.replace('T', ' ').split('.')[0] if ts else '—'
        
        r_vals = [
            f"REV-{r_idx:04d}",
            ts_clean,
            stall_code,
            stall_num_str,
            stall_name_str,
            stall_cat_str,
            star_num,
            star_display,
            r.get('review_text', '') or '—'
        ]
        
        ws3.row_dimensions[rev_data_row].height = 20
        is_even_rev = (rev_data_row % 2 == 0)
        
        for col_idx, val in enumerate(r_vals, 1):
            c = ws3.cell(row=rev_data_row, column=col_idx)
            c.value = val
            c.font = font_data
            c.border = border_cell
            c.alignment = rev_headers[col_idx-1][2]
            
            if is_even_rev:
                c.fill = fill_zebra
                
            if col_idx == 7:
                c.number_format = "#,##0"
                c.font = font_data_bold
            elif col_idx == 8:
                c.font = font_data_bold
                if star_num == 5:
                    c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                elif star_num == 4:
                    c.fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
                elif star_num <= 2:
                    c.fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
                    
        rev_data_row += 1

    # ----------------------------------------------------
    # SHEET 1: LEADERBOARD & STALL RATINGS
    # ----------------------------------------------------
    ws1 = wb.create_sheet(title="Stall Ratings Leaderboard", index=0)
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws1.merge_cells("A1:Q1")
    ws1["A1"] = "ADAMAS UNIVERSITY ENTREPRENEURSHIP CLUB (E-CLUB)"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_header_primary
    ws1["A1"].alignment = align_center
    ws1.row_dimensions[1].height = 28
    
    ws1.merge_cells("A2:Q2")
    ws1["A2"] = "RAKHI STARTUP BAZAAR — OFFICIAL STALL EVALUATION & VISITOR RATINGS REPORT"
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = PatternFill(start_color="6D0B14", end_color="6D0B14", fill_type="solid")
    ws1["A2"].alignment = align_center
    ws1.row_dimensions[2].height = 20

    # Summary KPI Cards (Rows 4-5) with Excel formulas
    ws1.merge_cells("B4:D4")
    ws1["B4"] = "TOTAL REGISTERED STALLS"
    ws1["B4"].font = font_kpi_lbl
    ws1["B4"].fill = fill_kpi_card
    ws1["B4"].alignment = align_center
    ws1.merge_cells("B5:D5")
    ws1["B5"] = "=COUNTA(B8:B39)"
    ws1["B5"].font = font_kpi_num
    ws1["B5"].fill = fill_kpi_card
    ws1["B5"].alignment = align_center
    ws1["B5"].number_format = '#,##0'

    ws1.merge_cells("F4:H4")
    ws1["F4"] = "TOTAL VISITOR RATINGS LOGGED"
    ws1["F4"].font = font_kpi_lbl
    ws1["F4"].fill = fill_kpi_card
    ws1["F4"].alignment = align_center
    ws1.merge_cells("F5:H5")
    ws1["F5"] = "=SUM(G8:G39)"
    ws1["F5"].font = font_kpi_num
    ws1["F5"].fill = fill_kpi_card
    ws1["F5"].alignment = align_center
    ws1["F5"].number_format = '#,##0'

    ws1.merge_cells("J4:L4")
    ws1["J4"] = "OVERALL EVENT AVG RATING"
    ws1["J4"].font = font_kpi_lbl
    ws1["J4"].fill = fill_kpi_card
    ws1["J4"].alignment = align_center
    ws1.merge_cells("J5:L5")
    ws1["J5"] = "=AVERAGE('All Visitor Reviews (1,178)'!G4:G" + str(last_review_row) + ")"
    ws1["J5"].font = font_kpi_num
    ws1["J5"].fill = fill_kpi_card
    ws1["J5"].alignment = align_center
    ws1["J5"].number_format = '0.00" ★"'

    ws1.merge_cells("N4:P4")
    ws1["N4"] = "EVENT 5★ EXCELLENCE VOTES"
    ws1["N4"].font = font_kpi_lbl
    ws1["N4"].fill = fill_kpi_card
    ws1["N4"].alignment = align_center
    ws1.merge_cells("N5:P5")
    ws1["N5"] = "=SUM(I8:I39)"
    ws1["N5"].font = font_kpi_num
    ws1["N5"].fill = fill_kpi_card
    ws1["N5"].alignment = align_center
    ws1["N5"].number_format = '#,##0'

    ws1.row_dimensions[4].height = 18
    ws1.row_dimensions[5].height = 26

    # Table Headers (Row 7)
    headers_ws1 = [
        ("Rank", 8, align_center),
        ("Stall #", 11, align_center),
        ("Stall / Venture Name", 28, align_left),
        ("Venture Category", 26, align_left),
        ("Student Founder(s)", 26, align_left),
        ("Department / School", 32, align_left),
        ("Total Votes", 12, align_center),
        ("Average Rating", 14, align_center),
        ("5★ (Excel)", 10, align_center),
        ("4★ (Good)", 10, align_center),
        ("3★ (Fair)", 10, align_center),
        ("2★ (Poor)", 10, align_center),
        ("1★ (Low)", 10, align_center),
        ("5★ Share (%)", 13, align_center),
        ("Top Visitor Highlights", 36, align_left),
        ("Founder Contact", 15, align_center),
        ("Instagram Handle", 22, align_left),
    ]

    header_row = 7
    ws1.row_dimensions[header_row].height = 26
    
    for col_idx, (h_title, col_width, col_align) in enumerate(headers_ws1, 1):
        c = ws1.cell(row=header_row, column=col_idx)
        c.value = h_title
        c.font = font_tbl_header
        c.fill = fill_header_primary
        c.alignment = align_center
        c.border = border_header
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = col_width

    # Data Rows (Rows 8 to 39 for 32 stalls)
    current_row = 8
    reviews_sheet_name = "'All Visitor Reviews (1,178)'"
    
    for idx, sid in enumerate(ranked_stall_ids, 1):
        item = stall_stats[sid]
        s = item['stall']
        s_num_padded = str(s.get('stall_number', '')).strip().zfill(2)
        
        # Excel Formulas linked to Sheet 3
        formula_total = f"=COUNTIF({reviews_sheet_name}!$C$4:$C${last_review_row}, \"{s_num_padded}\")"
        formula_5 = f"=COUNTIFS({reviews_sheet_name}!$C$4:$C${last_review_row}, \"{s_num_padded}\", {reviews_sheet_name}!$G$4:$G${last_review_row}, 5)"
        formula_4 = f"=COUNTIFS({reviews_sheet_name}!$C$4:$C${last_review_row}, \"{s_num_padded}\", {reviews_sheet_name}!$G$4:$G${last_review_row}, 4)"
        formula_3 = f"=COUNTIFS({reviews_sheet_name}!$C$4:$C${last_review_row}, \"{s_num_padded}\", {reviews_sheet_name}!$G$4:$G${last_review_row}, 3)"
        formula_2 = f"=COUNTIFS({reviews_sheet_name}!$C$4:$C${last_review_row}, \"{s_num_padded}\", {reviews_sheet_name}!$G$4:$G${last_review_row}, 2)"
        formula_1 = f"=COUNTIFS({reviews_sheet_name}!$C$4:$C${last_review_row}, \"{s_num_padded}\", {reviews_sheet_name}!$G$4:$G${last_review_row}, 1)"
        formula_avg = f"=IF(G{current_row}>0, (I{current_row}*5 + J{current_row}*4 + K{current_row}*3 + L{current_row}*2 + M{current_row}*1)/G{current_row}, 0)"
        formula_share = f"=IF(G{current_row}>0, I{current_row}/G{current_row}, 0)"
        
        row_data = [
            f"#{idx}",
            f"Stall #{s_num_padded}",
            s.get('name', ''),
            s.get('category', ''),
            s.get('founders', '—') or '—',
            s.get('department', '—') or '—',
            formula_total,
            formula_avg,
            formula_5,
            formula_4,
            formula_3,
            formula_2,
            formula_1,
            formula_share,
            item['top_tags'],
            s.get('contact', '—') or '—',
            s.get('instagram', '—') or '—'
        ]
        
        ws1.row_dimensions[current_row].height = 21
        is_even = (current_row % 2 == 0)
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=current_row, column=col_idx)
            cell.value = val
            cell.font = font_data
            cell.border = border_cell
            
            if is_even:
                cell.fill = fill_zebra
                
            if idx == 1:
                cell.fill = fill_top1
            elif idx == 2:
                cell.fill = fill_top2
            elif idx == 3:
                cell.fill = fill_top3
                
            if col_idx in [1, 2]:
                cell.alignment = align_center
                cell.font = font_data_bold
            elif col_idx in [3, 4, 5, 6, 17]:
                cell.alignment = align_left
            elif col_idx in [7, 9, 10, 11, 12, 13]:
                cell.alignment = align_center
                cell.number_format = "#,##0"
            elif col_idx == 8:
                cell.alignment = align_center
                cell.number_format = '0.00" ★"'
                cell.font = font_data_bold
            elif col_idx == 14:
                cell.alignment = align_center
                cell.number_format = "0.0%"
            elif col_idx in [15, 16]:
                cell.alignment = align_left

        current_row += 1

    # Total / Summary Row (Row 40)
    ws1.row_dimensions[current_row].height = 23
    total_cells = [
        "TOTAL",
        f"{len(stalls)} Stalls",
        "—",
        "—",
        "—",
        "—",
        f"=SUM(G8:G{current_row-1})",
        f"=AVERAGE(H8:H{current_row-1})",
        f"=SUM(I8:I{current_row-1})",
        f"=SUM(J8:J{current_row-1})",
        f"=SUM(K8:K{current_row-1})",
        f"=SUM(L8:L{current_row-1})",
        f"=SUM(M8:M{current_row-1})",
        f"=AVERAGE(N8:N{current_row-1})",
        "—",
        "—",
        "—"
    ]
    
    for col_idx, val in enumerate(total_cells, 1):
        cell = ws1.cell(row=current_row, column=col_idx)
        cell.value = val
        cell.font = font_data_bold
        cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        cell.border = border_total_row
        cell.alignment = align_center if col_idx not in [3, 4, 5, 6] else align_left
        
        if col_idx in [7, 9, 10, 11, 12, 13]:
            cell.number_format = "#,##0"
        elif col_idx == 8:
            cell.number_format = '0.00" ★"'
        elif col_idx == 14:
            cell.number_format = "0.0%"

    # ----------------------------------------------------
    # SHEET 2: CATEGORY PERFORMANCE ANALYSIS
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Category Analysis", index=1)
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:G1")
    ws2["A1"] = "RAKHI STARTUP BAZAAR — VENTURE SEGMENT & CATEGORY BENCHMARKS"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_header_primary
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 28
    
    cat_headers = [
        ("Venture Category Segment", 32, align_left),
        ("Stall Count", 14, align_center),
        ("Total Votes Logged", 18, align_center),
        ("Average Star Rating", 18, align_center),
        ("5★ Excellence Votes", 18, align_center),
        ("Top Rated Stall in Segment", 34, align_left),
        ("Top Segment Score", 20, align_center),
    ]
    
    header_row_cat = 3
    ws2.row_dimensions[header_row_cat].height = 24
    for col_idx, (h_title, col_width, col_align) in enumerate(cat_headers, 1):
        c = ws2.cell(row=header_row_cat, column=col_idx)
        c.value = h_title
        c.font = font_tbl_header
        c.fill = fill_header_secondary
        c.alignment = align_center
        c.border = border_header
        ws2.column_dimensions[get_column_letter(col_idx)].width = col_width
        
    categories = sorted(list(set([s.get('category', 'Uncategorized') for s in stalls])))
    cat_row = 4
    for cat in categories:
        cat_stalls = [s for s in stalls if s.get('category') == cat]
        
        # Top stall in this cat for display
        top_cat_stall = max(cat_stalls, key=lambda s: (stall_stats[s['id']]['avg_rating'], stall_stats[s['id']]['total_reviews']))
        top_cat_info = stall_stats[top_cat_stall['id']]
        top_cat_str = f"Stall #{str(top_cat_stall.get('stall_number','')).zfill(2)}: {top_cat_stall.get('name','')}"
        top_cat_score = f"{top_cat_info['avg_rating']:.2f} ★ ({top_cat_info['total_reviews']} votes)"
        
        # Formulas linking to Leaderboard sheet
        formula_stalls = f"=COUNTIF('Stall Ratings Leaderboard'!$D$8:$D$39, A{cat_row})"
        formula_votes = f"=SUMIF('Stall Ratings Leaderboard'!$D$8:$D$39, A{cat_row}, 'Stall Ratings Leaderboard'!$G$8:$G$39)"
        formula_avg_cat = f"=AVERAGEIF('Stall Ratings Leaderboard'!$D$8:$D$39, A{cat_row}, 'Stall Ratings Leaderboard'!$H$8:$H$39)"
        formula_5_stars = f"=SUMIF('Stall Ratings Leaderboard'!$D$8:$D$39, A{cat_row}, 'Stall Ratings Leaderboard'!$I$8:$I$39)"
        
        ws2.row_dimensions[cat_row].height = 22
        row_vals = [
            cat,
            formula_stalls,
            formula_votes,
            formula_avg_cat,
            formula_5_stars,
            top_cat_str,
            top_cat_score
        ]
        
        for col_idx, val in enumerate(row_vals, 1):
            c = ws2.cell(row=cat_row, column=col_idx)
            c.value = val
            c.font = font_data
            c.border = border_cell
            c.alignment = cat_headers[col_idx-1][2]
            
            if col_idx in [2, 3, 5]:
                c.number_format = "#,##0"
            elif col_idx == 4:
                c.number_format = '0.00" ★"'
                c.font = font_data_bold
                
        cat_row += 1

    # Total Row for Category sheet
    ws2.row_dimensions[cat_row].height = 23
    cat_total_vals = [
        "TOTAL",
        f"=SUM(B4:B{cat_row-1})",
        f"=SUM(C4:C{cat_row-1})",
        f"=AVERAGE(D4:D{cat_row-1})",
        f"=SUM(E4:E{cat_row-1})",
        "—",
        "—"
    ]
    for col_idx, val in enumerate(cat_total_vals, 1):
        c = ws2.cell(row=cat_row, column=col_idx)
        c.value = val
        c.font = font_data_bold
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        c.border = border_total_row
        c.alignment = cat_headers[col_idx-1][2]
        
        if col_idx in [2, 3, 5]:
            c.number_format = "#,##0"
        elif col_idx == 4:
            c.number_format = '0.00" ★"'

    output_filename = "Rakhi_Startup_Bazaar_Stall_Ratings_Report.xlsx"
    wb.save(output_filename)
    print(f"Successfully generated dynamic, professionally styled Excel file: {output_filename}")

if __name__ == '__main__':
    generate_excel_report()
